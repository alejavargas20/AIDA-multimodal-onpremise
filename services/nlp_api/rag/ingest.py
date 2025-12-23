import os
import hashlib
from pathlib import Path
from datetime import datetime, timezone
from typing import List, Optional

import pdfplumber
from docx import Document as DocxDocument

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.embeddings import HuggingFaceEmbeddings
from langchain_community.vectorstores import Chroma


DATA_RAW_DIR = os.getenv("DATA_RAW_DIR", os.getenv("RAW_DIR", "data/raw"))
CHROMA_DIR = os.getenv("CHROMA_DIR", "data/chroma")
EMBEDDING_MODEL = os.getenv(
    "EMBEDDING_MODEL",
    "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2",
)

CHUNK_SIZE = int(os.getenv("CHUNK_SIZE", "900"))
CHUNK_OVERLAP = int(os.getenv("CHUNK_OVERLAP", "150"))


# Loaders 
def load_txt(path: Path) -> List[Document]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return [Document(page_content=text, metadata={"source": str(path), "page": None})]


def load_docx(path: Path) -> List[Document]:
    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    text = "\n".join(parts)
    return [Document(page_content=text, metadata={"source": str(path), "page": None})]


def load_pdf(path: Path) -> List[Document]:
    """
    Sin OCR: extrae texto “tal cual” desde el PDF.
    Creamos 1 Document por página para poder guardar metadata.page.
    """
    docs: List[Document] = []
    with pdfplumber.open(str(path)) as pdf:
        for i, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            text = text.strip()
            if not text:
                continue
            docs.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": str(path),
                        "page": i + 1,  # 1-based para demo (más humano)
                    },
                )
            )
    return docs


def load_documents(raw_dir: Path) -> List[Document]:
    docs: List[Document] = []
    for p in sorted(raw_dir.rglob("*")):
        if not p.is_file():
            continue
        if p.name.startswith("."):
            continue

        ext = p.suffix.lower()
        try:
            if ext == ".txt":
                docs.extend(load_txt(p))
            elif ext == ".pdf":
                docs.extend(load_pdf(p))
            elif ext == ".docx":
                docs.extend(load_docx(p))
        except Exception as e:
            print(f"[WARN] No pude cargar {p} ({ext}): {e}")

    return docs


# Metadata por chunk 
def stable_doc_id(source: str, page: Optional[int]) -> str:
    base = f"{source}::page={page}"
    return hashlib.sha1(base.encode("utf-8")).hexdigest()[:12]


def enrich_chunk_metadata(chunks: List[Document]) -> List[Document]:
    now = datetime.now(timezone.utc).isoformat()
    for idx, d in enumerate(chunks):
        source = d.metadata.get("source")
        page = d.metadata.get("page", None)
        doc_id = stable_doc_id(str(source), page)

        # chunk_id: único y trazable
        d.metadata["chunk_id"] = f"{doc_id}:{idx}"
        d.metadata["created_at"] = now
        d.metadata["file_type"] = Path(str(source)).suffix.lower() if source else None

    return chunks


def ingest() -> None:
    raw_dir = Path(DATA_RAW_DIR)
    if not raw_dir.exists():
        raise RuntimeError(f"No existe la carpeta {raw_dir}")

    docs = load_documents(raw_dir)
    if not docs:
        print("No hay documentos en data/raw. Mete PDFs/TXTs/DOCX y vuelve a ejecutar.")
        return

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        separators=["\n\n", "\n", ". ", " ", ""],
    )

    chunks = splitter.split_documents(docs)
    chunks = enrich_chunk_metadata(chunks)

    emb = HuggingFaceEmbeddings(model_name=EMBEDDING_MODEL)

    _ = Chroma.from_documents(
        documents=chunks,
        embedding=emb,
        persist_directory=CHROMA_DIR,
    )

    print(f"OK: {len(chunks)} chunks indexados en {CHROMA_DIR}")
    print(f"Tipos soportados: .txt, .pdf, .docx | Embeddings: {EMBEDDING_MODEL}")


if __name__ == "__main__":
    ingest()
